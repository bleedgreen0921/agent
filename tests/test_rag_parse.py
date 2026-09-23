import sys
import types
from pathlib import Path

import httpx

from rag_service.chunking import make_chunks
from rag_service import models
from rag_service.parse import Block, parse_docling, parse_markdown, parse_txt


class CharacterTokenizer:
    def encode(self, text, add_special_tokens=False):
        return list(text)

    def __call__(self, text, add_special_tokens=False, return_offsets_mapping=False):
        return {"offset_mapping": [(i, i + 1) for i in range(len(text))]}


def test_txt_markdown_and_token_chunks():
    txt = parse_txt(b"one\n\ntwo")
    assert [item.locator["line_start"] for item in txt] == [1, 3]
    md = parse_markdown("# Heading\n\nFirst paragraph.\n\n| A | B |\n|---|---|\n| 1 | 2 |\n")
    assert md[0].heading_path == ("Heading",)
    assert any(item.kind == "table" and item.locator["line_start"] == 5 for item in md)
    blocks = [Block("paragraph", "x" * 100, {"kind": "txt", "line_start": 1, "line_end": 1})]
    chunks = make_chunks(blocks, "T", CharacterTokenizer(), target=70, overlap=10)
    assert len(chunks) > 1
    assert chunks[0].content[-10:] == chunks[1].content[:10]
    table = Block("table", "a" * 200, {"kind": "txt", "line_start": 2, "line_end": 4}, table_html="<table/>")
    assert make_chunks([table], "T", CharacterTokenizer(), target=70, overlap=10)[0].table_html == "<table/>"
    sections = [
        Block("paragraph", "first", {"kind": "txt", "line_start": 1, "line_end": 1}, ("A",)),
        Block("paragraph", "second", {"kind": "txt", "line_start": 2, "line_end": 2}, ("B",)),
    ]
    section_chunks = make_chunks(sections, "T", CharacterTokenizer(), target=70, overlap=10)
    assert [item.heading_path for item in section_chunks] == [("A",), ("B",)]


def test_docling_mapping_preserves_table_and_locator(monkeypatch):
    class Item:
        def __init__(self, label, text, page):
            self.label, self.text = label, text
            self.prov = [types.SimpleNamespace(page_no=page)] if page else []

        def export_to_html(self, doc):
            return "<table><tr><td>value</td></tr></table>"

    class Document:
        def iterate_items(self):
            yield Item("section_header", "Heading", 1), 1
            yield Item("text", "body", 2), 2
            yield Item("table", "value", 3), 2

    package = types.ModuleType("docling")
    module = types.ModuleType("docling.document_converter")
    module.DocumentConverter = lambda: types.SimpleNamespace(convert=lambda path: types.SimpleNamespace(document=Document()))
    monkeypatch.setitem(sys.modules, "docling", package)
    monkeypatch.setitem(sys.modules, "docling.document_converter", module)
    blocks, warnings = parse_docling(Path("mock.pdf"), "pdf")
    assert blocks[0].locator == {"kind": "pdf", "page_start": 2, "page_end": 2}
    assert blocks[1].table_html.startswith("<table>")
    assert warnings == []


def test_synthetic_docx_with_docling(tmp_path):
    from docx import Document

    path = tmp_path / "fixture.docx"
    document = Document()
    document.add_heading("Procedure", level=1)
    document.add_paragraph("Synthetic procedure text.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    document.save(path)
    blocks, warnings = parse_docling(path, "docx")
    assert blocks[0].locator["heading_path"] == ["Procedure"]
    assert blocks[1].kind == "table" and "<table" in blocks[1].table_html
    assert warnings == []


def test_model_http_contracts(monkeypatch):
    def embedding_handler(request: httpx.Request):
        assert request.url.path == "/v1/embeddings"
        assert request.read() == b'{"model":"embedding-model","input":["one","two"]}'
        return httpx.Response(200, json={"data": [{"index": 1, "embedding": [0.0, 1.0]}, {"index": 0, "embedding": [1.0, 0.0]}]})

    monkeypatch.setenv("RAG_EMBEDDING_URL", "http://model.test")
    monkeypatch.setattr(models, "_client", lambda: httpx.Client(transport=httpx.MockTransport(embedding_handler)))
    assert models.embed(["one", "two"], "embedding-model", 2) == [[1.0, 0.0], [0.0, 1.0]]

    def rerank_handler(request: httpx.Request):
        assert request.url.path == "/rerank"
        payload = __import__("json").loads(request.read())
        assert payload["query"] == "question" and payload["documents"] == ["low", "high"]
        return httpx.Response(200, json={"results": [{"index": 0, "relevance_score": 0.1}, {"index": 1, "relevance_score": 0.9}]})

    monkeypatch.setenv("RAG_RERANK_URL", "http://rerank.test")
    monkeypatch.setattr(models, "_client", lambda: httpx.Client(transport=httpx.MockTransport(rerank_handler)))
    assert models.rerank("question", ["low", "high"]) == [1, 0]
